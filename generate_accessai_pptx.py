from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor


OUT_PATH = r"c:\III DIPLOMA DE EXTENSIÓN UNIVERSITARIA EN INTELIGENCIA ARTIFICIAL AVANZADA SAMSUNG INNOVATION CAMPUS (2025-26)\MÓDULO 10. Proyectos\AccessAI_Capstone_Presentation_Branding.pptx"


def set_background(slide, color=RGBColor(245, 247, 250)):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title(slide, title, subtitle=None):
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12.0), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = RGBColor(16, 35, 72)
    p.alignment = PP_ALIGN.LEFT

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.0), Inches(11.5), Inches(0.4))
        tf2 = sub_box.text_frame
        p2 = tf2.paragraphs[0]
        run2 = p2.add_run()
        run2.text = subtitle
        run2.font.size = Pt(12)
        run2.font.color.rgb = RGBColor(90, 105, 125)
        p2.alignment = PP_ALIGN.LEFT


def add_bullets(slide, x, y, w, h, items, font_size=18, color=RGBColor(29, 41, 57), bullet_color=RGBColor(26, 115, 232)):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.bullet = True
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(8)
        for run in p.runs:
            run.font.size = Pt(font_size)
            run.font.color.rgb = color
            run.font.name = 'Aptos'
    return box


def add_card(slide, left, top, width, height, title, body, accent=RGBColor(26, 115, 232), title_color=RGBColor(16, 35, 72), body_color=RGBColor(55, 67, 84)):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = accent
    shape.line.width = Pt(1.5)

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()

    title_box = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(0.22), width - Inches(0.3), Inches(0.5))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = title_color

    body_box = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(0.7), width - Inches(0.3), height - Inches(0.9))
    tf2 = body_box.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = body
    p2.alignment = PP_ALIGN.LEFT
    for run in p2.runs:
        run.font.size = Pt(12)
        run.font.color.rgb = body_color
        run.font.name = 'Aptos'


def add_flow_box(slide, left, top, width, height, text, fill_color, text_color=RGBColor(255,255,255)):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = fill_color
    box.line.color.rgb = fill_color
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    for run in p.runs:
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = text_color
        run.font.name = 'Aptos'
    return box


def add_arrow(slide, x1, y1, x2, y2):
    arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x1, y1, x2 - x1, y2 - y1)
    arrow.fill.solid(); arrow.fill.fore_color.rgb = RGBColor(61, 92, 125)
    arrow.line.color.rgb = RGBColor(61, 92, 125)
    return arrow


def add_brand_block(slide, left, top, width, height, text, fill_color, font_size=12):
    block = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    block.fill.solid()
    block.fill.fore_color.rgb = fill_color
    block.line.color.rgb = fill_color
    tf = block.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    for run in p.runs:
        run.font.size = Pt(font_size)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.name = 'Aptos'
    return block


def add_logo_marker(slide, left, top, width, height, label, fill_color, text_color=RGBColor(255,255,255), font_size=12):
    marker = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    marker.fill.solid()
    marker.fill.fore_color.rgb = fill_color
    marker.line.color.rgb = fill_color
    tf = marker.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = label
    p.alignment = PP_ALIGN.CENTER
    for run in p.runs:
        run.font.size = Pt(font_size)
        run.font.bold = True
        run.font.color.rgb = text_color
        run.font.name = 'Aptos'
    return marker


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Slide 1: Title
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, RGBColor(240, 245, 250))

header_band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.28))
header_band.fill.solid(); header_band.fill.fore_color.rgb = RGBColor(20, 40, 160)
header_band.line.fill.background()

# institutional branding strip in Samsung palette
add_logo_marker(slide, Inches(0.7), Inches(0.45), Inches(1.3), Inches(0.56), 'ONCE', RGBColor(220, 0, 40), font_size=11)
add_logo_marker(slide, Inches(2.3), Inches(0.45), Inches(3.0), Inches(0.56), 'UNIVERSIDAD\nDE MÁLAGA', RGBColor(0, 87, 156), font_size=9)
add_logo_marker(slide, Inches(5.7), Inches(0.45), Inches(2.1), Inches(0.56), 'SAMSUNG', RGBColor(0, 33, 128), font_size=11)

team = slide.shapes.add_textbox(Inches(0.7), Inches(1.15), Inches(5.2), Inches(0.5))
tf = team.text_frame
p = tf.paragraphs[0]
p.text = 'Samsung Innovation Campus | Capstone Project'
for run in p.runs:
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 33, 128)

title = slide.shapes.add_textbox(Inches(0.7), Inches(1.7), Inches(11.5), Inches(1.3))
tf2 = title.text_frame
p2 = tf2.paragraphs[0]
p2.text = 'AccessAI\nAI-Based Urban Accessibility Detection'
for run in p2.runs:
    run.font.size = Pt(26 if len(p2.text) < 60 else 22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 35, 72)

subtitle = slide.shapes.add_textbox(Inches(0.7), Inches(3.05), Inches(8.5), Inches(0.7))
tf3 = subtitle.text_frame
p3 = tf3.paragraphs[0]
p3.text = 'Vision AI prototype for identifying sidewalks, ramps, and urban accessibility barriers.'
for run in p3.runs:
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(64, 74, 90)

# accent shapes with Samsung palette
accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.3), Inches(1.8), Inches(2.9), Inches(2.4))
accent.fill.solid(); accent.fill.fore_color.rgb = RGBColor(0, 33, 128)
accent.line.fill.background()

mini = slide.shapes.add_textbox(Inches(9.6), Inches(2.15), Inches(2.3), Inches(1.1))
tf4 = mini.text_frame
p4 = tf4.paragraphs[0]
p4.text = 'YOLOv8n\n+ Vision AI'
for run in p4.runs:
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255,255,255)

footer = slide.shapes.add_textbox(Inches(0.7), Inches(6.7), Inches(9.5), Inches(0.4))
tf5 = footer.text_frame
p5 = tf5.paragraphs[0]
p5.text = 'ONCE  |  Universidad de Málaga  |  Samsung Innovation Campus 2025-26'
for run in p5.runs:
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(85, 94, 109)

# Slide 2: Problem and objective
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, RGBColor(250, 251, 253))
add_title(slide, '1. Problem and Objective', 'Urban accessibility is often hidden behind seemingly accessible maps and streets.')

add_card(slide, Inches(0.7), Inches(1.5), Inches(5.4), Inches(4.2),
    'Background',
    'Many sidewalks appear accessible on paper maps, but people with reduced mobility may face barriers such as missing ramps, uneven surfaces, curbs, and obstacles in public spaces.',
    accent=RGBColor(37, 127, 223))

add_card(slide, Inches(6.5), Inches(1.5), Inches(6.1), Inches(4.2),
    'Objective',
    'Develop an AI prototype that detects accessibility-related elements in urban images and supports future navigation tools for wheelchair users, pedestrians, and urban planners.',
    accent=RGBColor(23, 164, 111))

add_bullets(slide, Inches(0.9), Inches(5.9), Inches(11.5), Inches(1.0), [
    'Goal: detect sidewalks and access ramps in initial prototype.',
    'Future scope: obstacles, curbs, deteriorated surfaces, stairs, and geolocated urban accessibility maps.'
], font_size=14)

# Slide 3: Data and methodology
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, RGBColor(245, 248, 252))
add_title(slide, '2. Data and Methodology', 'Public urban data and computer vision applied to a real-world accessibility problem.')

add_card(slide, Inches(0.7), Inches(1.5), Inches(5.9), Inches(4.8),
    'Data sources',
    'Project Sidewalk\nSidewalk Accessibility\nCityscapes\nPublic dataset review for image quality, class balance, annotation quality, and object density.',
    accent=RGBColor(26, 115, 232))

add_card(slide, Inches(6.8), Inches(1.5), Inches(5.8), Inches(4.8),
    'Methodology',
    'Supervised learning with computer vision.\n1) Data cleaning and split.\n2) Training and validation.\n3) YOLOv8n for object detection.\n4) Evaluation with precision, recall, F1 and mAP.\n5) Prototype deployment with image input and annotated output.',
    accent=RGBColor(88, 95, 255))

# Slide 4: Workflow
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, RGBColor(255, 255, 255))
add_title(slide, '3. Workflow and System Design', 'From data collection to image-based accessibility detection.')

add_flow_box(slide, Inches(0.7), Inches(2.2), Inches(2.2), Inches(1.25), 'Data\nCollection', RGBColor(30, 52, 110))
add_flow_box(slide, Inches(3.3), Inches(2.2), Inches(2.2), Inches(1.25), 'Data\nPreparation', RGBColor(26, 115, 232))
add_flow_box(slide, Inches(5.9), Inches(2.2), Inches(2.2), Inches(1.25), 'Model\nTraining', RGBColor(48, 154, 103))
add_flow_box(slide, Inches(8.5), Inches(2.2), Inches(2.1), Inches(1.25), 'Evaluation', RGBColor(159, 95, 255))
add_flow_box(slide, Inches(10.9), Inches(2.2), Inches(1.8), Inches(1.25), 'Prototype', RGBColor(232, 97, 76))

add_arrow(slide, Inches(2.9), Inches(2.8), Inches(3.3), Inches(2.8))
add_arrow(slide, Inches(5.5), Inches(2.8), Inches(5.9), Inches(2.8))
add_arrow(slide, Inches(8.1), Inches(2.8), Inches(8.5), Inches(2.8))
add_arrow(slide, Inches(10.6), Inches(2.8), Inches(10.9), Inches(2.8))

add_card(slide, Inches(1.1), Inches(4.2), Inches(11.1), Inches(2.2),
    'System design',
    'Input image -> cleaning and labeling -> YOLO object detection -> confidence threshold -> annotated image with accessibility classes and bounding boxes -> possible future integration with urban maps and accessible-route analysis.',
    accent=RGBColor(80, 125, 200))

# Slide 5: Results and demo
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, RGBColor(246, 250, 248))
add_title(slide, '4. Results and Demo', 'Prototype ready to process urban images and highlight detected accessibility elements.')

add_card(slide, Inches(0.7), Inches(1.5), Inches(4.0), Inches(4.6),
    'Current prototype',
    'The project includes a Python demo script that loads a YOLO model and processes an input image to identify accessible urban elements. The output is an annotated image with detected bounding boxes and class labels.',
    accent=RGBColor(25, 133, 85))

add_card(slide, Inches(4.9), Inches(1.5), Inches(3.8), Inches(4.6),
    'Key findings',
    'Initial focus: sidewalks and access ramps.\nGoal: detect the main urban accessibility barriers automatically.\nPotential expansion to stairs, obstacles, and poor surface conditions.',
    accent=RGBColor(44, 123, 207))

add_card(slide, Inches(8.9), Inches(1.5), Inches(3.8), Inches(4.6),
    'Demo path',
    'Example command: python accessai_demo.py --image imagen.jpg\nModel output saved as an annotated image for visualization and presentation.',
    accent=RGBColor(181, 112, 27))

# Slide 6: Impact and Next Steps
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, RGBColor(248, 244, 251))
add_title(slide, '5. Impact and Future Improvements', 'A practical AI tool with inclusive, social and urban value.')

add_card(slide, Inches(0.8), Inches(1.5), Inches(3.9), Inches(4.0),
    'Accomplishments',
    'Establishes a working computer vision project proposal for urban accessibility analysis. Produces a practical AI prototype and a clear roadmap for future improvements.',
    accent=RGBColor(94, 86, 255))

add_card(slide, Inches(4.95), Inches(1.5), Inches(3.9), Inches(4.0),
    'Benefits',
    'Helps support people with reduced mobility. Contributes to safer public spaces and better accessibility planning. Provides a foundation for route recommendations.',
    accent=RGBColor(34, 126, 110))

add_card(slide, Inches(9.1), Inches(1.5), Inches(3.3), Inches(4.0),
    'Next steps',
    'Expand dataset and labels. Improve detection of obstacles and ramps. Add georeferenced mapping and a user-facing mobile or web interface.',
    accent=RGBColor(220, 113, 44))

prs.save(OUT_PATH)
print(f'PowerPoint created: {OUT_PATH}')
