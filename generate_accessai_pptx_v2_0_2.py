from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

OUT_PATH = r"c:\III DIPLOMA DE EXTENSIÓN UNIVERSITARIA EN INTELIGENCIA ARTIFICIAL AVANZADA SAMSUNG INNOVATION CAMPUS (2025-26)\MÓDULO 10. Proyectos\AccessAI_Capstone_Presentation_v2_0_2.pptx"

# Samsung theme extracted from Theme3.thmx
THEME = {
    "dk1": RGBColor(0, 0, 0),
    "lt1": RGBColor(255, 255, 255),
    "dk2": RGBColor(68, 84, 106),
    "lt2": RGBColor(231, 230, 230),
    "accent1": RGBColor(91, 155, 213),
    "accent2": RGBColor(237, 125, 49),
    "accent3": RGBColor(165, 165, 165),
    "accent4": RGBColor(255, 192, 0),
    "accent5": RGBColor(68, 114, 196),
    "accent6": RGBColor(112, 173, 71),
    "hlink": RGBColor(5, 99, 193),
    "folHlink": RGBColor(149, 79, 114),
}


def set_bg(slide, color=THEME["lt1"]):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(slide, left, top, width, height, text, size=20, bold=False, color=THEME["dk2"], align=PP_ALIGN.LEFT, font_name='Samsung Sharp Sans'):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    p.text = text
    for run in p.runs:
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = font_name
        run.font.color.rgb = color
    return box


def add_bullets(slide, left, top, width, height, items, font_size=18, color=THEME["dk2"]):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.bullet = True
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(10)
        for run in p.runs:
            run.font.size = Pt(font_size)
            run.font.color.rgb = color
            run.font.name = 'Samsung Sharp Sans'
    return box


def add_card(slide, left, top, width, height, title, body, accent=THEME["accent1"]):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = THEME["lt1"]
    card.line.color.rgb = accent
    card.line.width = Pt(1.5)

    strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Inches(0.12))
    strip.fill.solid()
    strip.fill.fore_color.rgb = accent
    strip.line.color.rgb = accent

    add_text(slide, left + Inches(0.18), top + Inches(0.22), width - Inches(0.36), Inches(0.5), title, size=18, bold=True, color=THEME["dk2"])
    box = slide.shapes.add_textbox(left + Inches(0.18), top + Inches(0.7), width - Inches(0.36), height - Inches(0.9))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = body
    for run in p.runs:
        run.font.size = Pt(12)
        run.font.name = 'Samsung One 400'
        run.font.color.rgb = THEME["dk2"]


def add_step(slide, left, top, width, height, label, color=THEME["accent5"]):
    step = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    step.fill.solid()
    step.fill.fore_color.rgb = color
    step.line.color.rgb = color
    tf = step.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = label
    p.alignment = PP_ALIGN.CENTER
    for run in p.runs:
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.name = 'Samsung Sharp Sans'
        run.font.color.rgb = THEME["lt1"]
    return step


def add_arrow(slide, x1, y1, x2, y2):
    arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x1, y1, x2 - x1, y2 - y1)
    arrow.fill.solid(); arrow.fill.fore_color.rgb = THEME["accent3"]
    arrow.line.color.rgb = THEME["accent3"]
    return arrow


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Slide 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, THEME["lt2"])

band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.28))
band.fill.solid(); band.fill.fore_color.rgb = THEME["dk2"]
band.line.color.rgb = THEME["dk2"]

add_text(slide, Inches(0.6), Inches(0.45), Inches(2.2), Inches(0.45), 'ONCE', size=12, bold=True, color=THEME["lt1"])
add_text(slide, Inches(3.2), Inches(0.45), Inches(3.1), Inches(0.45), 'UNIVERSIDAD DE MÁLAGA', size=10, bold=True, color=THEME["lt1"])
add_text(slide, Inches(7.4), Inches(0.45), Inches(2.2), Inches(0.45), 'SAMSUNG', size=12, bold=True, color=THEME["lt1"])

add_text(slide, Inches(0.7), Inches(1.25), Inches(6.8), Inches(0.4), 'Samsung Innovation Campus | Capstone Project', size=14, bold=True, color=THEME["accent5"])
add_text(slide, Inches(0.7), Inches(1.75), Inches(9.5), Inches(1.05), 'AccessAI\nAI-Based Urban Accessibility Detection', size=28, bold=True, color=THEME["dk2"])
add_text(slide, Inches(0.7), Inches(3.1), Inches(8.0), Inches(0.5), 'Vision AI prototype for detecting sidewalks, ramps and urban accessibility barriers.', size=18, color=THEME["dk2"])

card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.4), Inches(1.7), Inches(2.8), Inches(2.3))
card.fill.solid(); card.fill.fore_color.rgb = THEME["accent1"]
card.line.color.rgb = THEME["accent1"]
add_text(slide, Inches(9.7), Inches(2.25), Inches(2.2), Inches(1.1), 'YOLOv8n\n+ Vision AI', size=18, bold=True, color=THEME["lt1"], align=PP_ALIGN.CENTER)

add_text(slide, Inches(0.7), Inches(6.7), Inches(9.0), Inches(0.35), 'ONCE | Universidad de Málaga | Samsung Innovation Campus 2025-26', size=10, color=THEME["dk2"])

# Slide 2
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, THEME["lt1"])
add_text(slide, Inches(0.6), Inches(0.4), Inches(8.0), Inches(0.7), '1. Problem and Objective', size=26, bold=True, color=THEME["dk2"])
add_text(slide, Inches(0.7), Inches(1.0), Inches(12.0), Inches(0.4), 'Urban accessibility is often hidden behind maps that look accessible but still include barriers.', size=13, color=THEME["dk2"])

add_card(slide, Inches(0.7), Inches(1.6), Inches(5.7), Inches(3.7), 'Background', 'Many sidewalks appear accessible in maps, but people with reduced mobility may face barriers such as missing ramps, uneven surfaces, elevated curbs, or obstacles in public spaces.')
add_card(slide, Inches(6.7), Inches(1.6), Inches(5.8), Inches(3.7), 'Objective', 'Develop an AI prototype to detect accessibility-related elements in urban images and support future tools for safer and easier navigation.')

add_bullets(slide, Inches(0.9), Inches(5.6), Inches(11.5), Inches(1.0), [
    'Initial scope: sidewalks and access ramps.',
    'Future extension: curbs, obstacles, stairs and geolocated accessibility maps.'
], font_size=14)

# Slide 3
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, THEME["lt1"])
add_text(slide, Inches(0.6), Inches(0.4), Inches(8.0), Inches(0.7), '2. Data and Methodology', size=26, bold=True, color=THEME["dk2"])

add_card(slide, Inches(0.7), Inches(1.5), Inches(5.6), Inches(4.3), 'Data sources', 'Project Sidewalk\nSidewalk Accessibility\nCityscapes\nReview of image quality, class balance, annotation quality, and object density.')
add_card(slide, Inches(6.7), Inches(1.5), Inches(5.8), Inches(4.3), 'Methodology', 'Supervised learning with computer vision.\n1) Data cleaning and split.\n2) Model training and validation.\n3) YOLOv8n object detection.\n4) Metrics: precision, recall, F1, mAP.\n5) Annotated prototype output.')

# Slide 4
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, THEME["lt1"])
add_text(slide, Inches(0.6), Inches(0.4), Inches(8.0), Inches(0.7), '3. Workflow and System Design', size=26, bold=True, color=THEME["dk2"])

add_step(slide, Inches(0.7), Inches(2.3), Inches(2.0), Inches(1.2), 'Data\nCollection', THEME["dk2"])
add_step(slide, Inches(3.0), Inches(2.3), Inches(2.2), Inches(1.2), 'Data\nPreparation', THEME["accent1"])
add_step(slide, Inches(5.6), Inches(2.3), Inches(2.0), Inches(1.2), 'Model\nTraining', THEME["accent6"])
add_step(slide, Inches(8.1), Inches(2.3), Inches(1.8), Inches(1.2), 'Evaluation', THEME["accent2"])
add_step(slide, Inches(10.5), Inches(2.3), Inches(1.9), Inches(1.2), 'Prototype', THEME["accent4"])

add_arrow(slide, Inches(2.7), Inches(2.9), Inches(3.0), Inches(2.9))
add_arrow(slide, Inches(5.2), Inches(2.9), Inches(5.6), Inches(2.9))
add_arrow(slide, Inches(7.6), Inches(2.9), Inches(8.1), Inches(2.9))
add_arrow(slide, Inches(9.9), Inches(2.9), Inches(10.5), Inches(2.9))

add_card(slide, Inches(1.0), Inches(4.3), Inches(11.2), Inches(2.0), 'System design', 'Input image -> cleaning and labeling -> YOLO detection -> confidence threshold -> annotated output with accessibility classes and bounding boxes -> future mapping support.')

# Slide 5
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, THEME["lt1"])
add_text(slide, Inches(0.6), Inches(0.4), Inches(8.0), Inches(0.7), '4. Results and Demo', size=26, bold=True, color=THEME["dk2"])

add_card(slide, Inches(0.7), Inches(1.5), Inches(3.8), Inches(4.2), 'Current prototype', 'The project includes a Python demo that loads a YOLO model and processes an input image to identify accessibility-related urban elements.')
add_card(slide, Inches(4.9), Inches(1.5), Inches(3.7), Inches(4.2), 'Key findings', 'Main focus: sidewalks and access ramps.\nPotential expansion: stairs, obstacles, and poor surface conditions.')
add_card(slide, Inches(8.9), Inches(1.5), Inches(3.6), Inches(4.2), 'Demo path', 'Command example: python accessai_demo.py --image imagen.jpg\nOutput is an annotated image for visualization and presentation.')

# Slide 6
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, THEME["lt2"])
add_text(slide, Inches(0.6), Inches(0.4), Inches(8.0), Inches(0.7), '5. Impact and Future Improvements', size=26, bold=True, color=THEME["dk2"])
add_card(slide, Inches(0.8), Inches(1.6), Inches(3.8), Inches(4.2), 'Accomplishments', 'The project defines a practical research and prototype path for urban accessibility detection using vision AI.')
add_card(slide, Inches(4.9), Inches(1.6), Inches(3.7), Inches(4.2), 'Benefits', 'Supports safer public spaces and helps people with reduced mobility by identifying accessibility barriers more efficiently.')
add_card(slide, Inches(9.0), Inches(1.6), Inches(3.5), Inches(4.2), 'Next steps', 'Expand datasets and labels. Improve detection of obstacles and ramps. Add geographic awareness and a user-facing interface.')

prs.save(OUT_PATH)
print(f'PowerPoint created: {OUT_PATH}')
